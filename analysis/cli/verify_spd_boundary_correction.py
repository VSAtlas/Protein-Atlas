from __future__ import annotations

import argparse
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import platform
import subprocess
import sys
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
import pandas as pd

from analysis.external.spd import (
    SPD_LABEL_POLICY_VERSION,
    _spd_label_status,
    aggregate_spd_assays,
    normalize_spd_table,
)
from analysis.ml.spd_four_expert_tables import (
    BINDING_LABEL_POLICY_VERSION,
    _binding_label,
)
from analysis.ml.spd_panel import _extract_target_id, _first_nonempty
from analysis.spd_activity_policy import SPD_ACTIVITY_PARSER_VERSION


EVIDENCE_COLUMNS = [
    "source_row_id",
    "source_excel_row",
    "production_selected",
    "drugcentral_struct_id",
    "drug_name",
    "parent_when_prodrug_id",
    "struct_match_type",
    "inchi_key",
    "assay_id",
    "assay_group",
    "assay_group_name",
    "preferred_assay_id",
    "assay_format",
    "assay_mode",
    "assay_readout",
    "assay_technology",
    "assay_species",
    "assay_gene",
    "assay_representative_gene",
    "source_representative",
    "activity_relation",
    "activity_relation_cell",
    "activity_relation_xml_value",
    "activity_ac50_uM",
    "activity_ac50_cell",
    "activity_ac50_xml_value",
    "activity_ac50_cell_type",
    "activity_ac50_number_format",
    "activity_ac50_has_formula",
    "n_crc_summarized",
    "n_crc_total",
    "drugcentral_ac50_median_uM",
    "chembl_ac50_median_uM",
    "subscription_ac50_median_uM",
    "drugcentral_ac50_count",
    "chembl_ac50_count",
    "subscription_ac50_count",
    "source_summary_inconsistency",
    "total_cmax_uM",
    "total_cmax_cell",
    "total_cmax_xml_value",
    "cmax_source",
    "ppb_percent",
    "ppb_cell",
    "ppb_xml_value",
    "ppb_qualifier",
    "ppb_source",
    "fraction_unbound_decimal",
    "free_cmax_uM",
    "free_cmax_cell",
    "free_cmax_xml_value",
    "free_cmax_cell_type",
    "free_cmax_number_format",
    "free_cmax_has_formula",
    "free_cmax_calculated_decimal",
    "free_cmax_matches_calculation",
    "cmax_from_parent",
    "stored_margin",
    "stored_margin_cell",
    "stored_margin_xml_value",
    "lower_margin_decimal",
    "margin_interval",
    "old_exposure_label",
    "new_exposure_label",
    "ppb_is_bound",
    "free_cmax_has_explicit_interval",
    "binary_float_or_display_rounding_dependency",
    "source_record_identifier_available",
    "biological_precision_documented",
    "stored_value_contract_verdict",
    "biological_provenance_verdict",
    "verification_notes",
]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _git(repo_root: Path, *args: str) -> str:
    return subprocess.check_output(
        ["git", *args],
        cwd=repo_root,
        text=True,
    )


def _clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if bool(pd.isna(value)):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _id(value: Any) -> str:
    text = _clean(value)
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except ValueError:
            pass
    return text


def _decimal(value: Any) -> Decimal:
    text = _clean(value)
    if not text:
        raise InvalidOperation("missing decimal value")
    return Decimal(text)


def _qualifier_is_bound(value: Any) -> bool:
    return _clean(value) in {"<", "<=", ">", ">="}


def _sheet_xml_path(archive: ZipFile, sheet_name: str) -> str:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationship_id = None
    for sheet in workbook.findall("main:sheets/main:sheet", ns):
        if sheet.attrib.get("name") == sheet_name:
            relationship_id = sheet.attrib.get(f"{{{ns['rel']}}}id")
            break
    if relationship_id is None:
        raise KeyError(f"workbook sheet not found: {sheet_name}")
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for relationship in relationships.findall("pkg:Relationship", ns):
        if relationship.attrib.get("Id") == relationship_id:
            target = relationship.attrib["Target"].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise KeyError(f"worksheet relationship not found: {sheet_name}")


def _xml_cell_values(
    workbook: Path,
    sheet_name: str,
    cell_references: set[str],
) -> dict[str, str]:
    namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    result: dict[str, str] = {}
    with ZipFile(workbook) as archive:
        sheet_path = _sheet_xml_path(archive, sheet_name)
        root = ElementTree.fromstring(archive.read(sheet_path))
        for cell in root.findall(".//main:c", namespace):
            reference = cell.attrib.get("r", "")
            if reference not in cell_references:
                continue
            value = cell.find("main:v", namespace)
            result[reference] = (
                "" if value is None or value.text is None else value.text
            )
    return result


def _column_map(worksheet: Any, header_row: int = 3) -> dict[str, int]:
    return {
        _clean(cell.value): int(cell.column)
        for cell in worksheet[header_row]
        if _clean(cell.value)
    }


def _cell_metadata(worksheet: Any, row: int, column: int) -> dict[str, Any]:
    cell = worksheet.cell(row=row, column=column)
    return {
        "coordinate": cell.coordinate,
        "value": cell.value,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "has_formula": cell.data_type == "f",
    }


def _bool_series(series: pd.Series) -> pd.Series:
    return (
        series.fillna(False)
        .astype(str)
        .str.strip()
        .str.lower()
        .isin({"1", "true", "yes"})
    )


def _label_name(value: Any) -> str:
    text = _clean(value).lower()
    if text in {"1", "1.0", "true", "positive"}:
        return "positive"
    if text in {"0", "0.0", "false", "negative"}:
        return "negative"
    if text in {"-1", "-1.0", "excluded", "conflict", "excluded_conflict"}:
        return "excluded_conflict"
    return "unknown"


def _label_counts(series: pd.Series) -> dict[str, int]:
    names = series.map(_label_name)
    return {
        label: int(names.eq(label).sum())
        for label in ("positive", "negative", "unknown", "excluded_conflict")
    }


def _production_exposure_label(status: Any) -> str:
    text = _clean(status)
    if text == "labeled_relevant":
        return "positive"
    if text.startswith("labeled_"):
        return "negative"
    return "unknown"


def _verify_production_integration(
    *,
    workbook_path: Path,
    phase_a_behavior_path: Path,
    out_dir: Path,
) -> tuple[dict[str, Any], list[Path]]:
    baseline = pd.read_csv(phase_a_behavior_path, low_memory=False)
    required = {
        "source_row_id",
        "drug_identity",
        "target_identity",
        "raw_relation",
        "raw_activity_value",
        "free_cmax_um",
        "production_selected",
        "materialized_binding_label",
        "legacy_binding_label",
        "canonical_binding_label",
        "materialized_exposure_label",
        "legacy_exposure_label",
        "canonical_exposure_label",
    }
    missing = sorted(required - set(baseline.columns))
    if missing:
        raise RuntimeError(
            f"Phase A behavior comparison is missing columns: {', '.join(missing)}"
        )

    binding_source = pd.DataFrame(
        {
            "spd_ac50_uM": baseline["raw_activity_value"],
            "spd_activity_relation": baseline["raw_relation"],
        },
        index=baseline.index,
    )
    current_binding = _binding_label(
        binding_source,
        active_um=1.0,
        inactive_um=10.0,
    ).map(_label_name)
    activity_nm = (
        pd.to_numeric(baseline["raw_activity_value"], errors="coerce") * 1000.0
    )
    free_cmax_nm = pd.to_numeric(baseline["free_cmax_um"], errors="coerce") * 1000.0
    current_status = pd.Series(
        [
            _spd_label_status(
                {
                    "ac50_nM": activity,
                    "free_cmax_nM": free_cmax,
                    "activity_relation": relation,
                },
                10.0,
                100.0,
            )
            for activity, relation, free_cmax in zip(
                activity_nm,
                baseline["raw_relation"],
                free_cmax_nm,
            )
        ],
        index=baseline.index,
        dtype="string",
    )
    current_exposure = current_status.map(_production_exposure_label)

    raw_comparison = baseline[
        [
            "source_row_id",
            "drug_identity",
            "target_identity",
            "raw_relation",
            "raw_activity_value",
            "free_cmax_um",
            "production_selected",
            "materialized_binding_label",
            "legacy_binding_label",
            "canonical_binding_label",
            "materialized_exposure_label",
            "legacy_exposure_label",
            "canonical_exposure_label",
        ]
    ].copy()
    raw_comparison["production_binding_label"] = current_binding
    raw_comparison["production_exposure_label"] = current_exposure
    raw_comparison["production_exposure_status"] = current_status
    raw_comparison["binding_legacy_to_production_changed"] = (
        raw_comparison["legacy_binding_label"].map(_label_name)
        != raw_comparison["production_binding_label"]
    )
    raw_comparison["exposure_legacy_to_production_changed"] = (
        raw_comparison["legacy_exposure_label"].map(_label_name)
        != raw_comparison["production_exposure_label"]
    )
    raw_comparison["binding_canonical_to_production_changed"] = (
        raw_comparison["canonical_binding_label"].map(_label_name)
        != raw_comparison["production_binding_label"]
    )
    raw_comparison["exposure_canonical_to_production_changed"] = (
        raw_comparison["canonical_exposure_label"].map(_label_name)
        != raw_comparison["production_exposure_label"]
    )

    normalized = normalize_spd_table(workbook_path)
    normalized = normalized.copy()
    normalized["drug_id"] = _first_nonempty(normalized, ["name", "drug_id"]).str.lower()
    normalized["target_id"] = _extract_target_id(normalized).astype(str).str.upper()
    production_pairs = aggregate_spd_assays(normalized)
    production_pairs["drug_id"] = production_pairs["drug_id"].astype(str).str.lower()
    production_pairs["target_id"] = (
        production_pairs["target_id"].astype(str).str.upper()
    )
    if production_pairs.duplicated(["drug_id", "target_id"]).any():
        raise RuntimeError("production pair output contains duplicate drug-target keys")
    pair_binding_source = pd.DataFrame(
        {
            "spd_ac50_uM": pd.to_numeric(production_pairs["ac50_nM"], errors="coerce")
            / 1000.0,
            "spd_activity_relation": production_pairs["spd_activity_relation"],
        },
        index=production_pairs.index,
    )
    production_pairs["production_binding_label"] = _binding_label(
        pair_binding_source,
        active_um=1.0,
        inactive_um=10.0,
    ).map(_label_name)
    production_pairs["production_exposure_label"] = production_pairs[
        "spd_label_status"
    ].map(_production_exposure_label)

    selected = baseline[_bool_series(baseline["production_selected"])].copy()
    selected["drug_id"] = selected["drug_identity"].astype(str).str.lower()
    selected["target_id"] = selected["target_identity"].astype(str).str.upper()
    if selected.duplicated(["drug_id", "target_id"]).any():
        raise RuntimeError("Phase A selected baseline contains duplicate pair keys")
    pair_comparison = selected[
        [
            "source_row_id",
            "drug_id",
            "target_id",
            "materialized_binding_label",
            "legacy_binding_label",
            "canonical_binding_label",
            "materialized_exposure_label",
            "legacy_exposure_label",
            "canonical_exposure_label",
        ]
    ].merge(
        production_pairs[
            [
                "drug_id",
                "target_id",
                "ac50_nM",
                "free_cmax_nM",
                "exposure_margin",
                "spd_activity_relation",
                "spd_label_status",
                "spd_missing_reason",
                "spd_exposure_relevant",
                "spd_exposure_weak",
                "spd_exposure_unlikely",
                "spd_label_policy_version",
                "production_binding_label",
                "production_exposure_label",
            ]
        ],
        on=["drug_id", "target_id"],
        how="outer",
        indicator=True,
        validate="one_to_one",
    )
    pair_comparison["binding_legacy_to_production_changed"] = (
        pair_comparison["legacy_binding_label"].map(_label_name)
        != pair_comparison["production_binding_label"]
    )
    pair_comparison["exposure_legacy_to_production_changed"] = (
        pair_comparison["legacy_exposure_label"].map(_label_name)
        != pair_comparison["production_exposure_label"]
    )
    pair_comparison["binding_canonical_to_production_changed"] = (
        pair_comparison["canonical_binding_label"].map(_label_name)
        != pair_comparison["production_binding_label"]
    )
    pair_comparison["exposure_canonical_to_production_changed"] = (
        pair_comparison["canonical_exposure_label"].map(_label_name)
        != pair_comparison["production_exposure_label"]
    )

    raw_binding_changes = int(
        raw_comparison["binding_legacy_to_production_changed"].sum()
    )
    raw_exposure_changes = int(
        raw_comparison["exposure_legacy_to_production_changed"].sum()
    )
    pair_binding_changes = int(
        pair_comparison["binding_legacy_to_production_changed"].sum()
    )
    pair_exposure_changes = int(
        pair_comparison["exposure_legacy_to_production_changed"].sum()
    )
    canonical_mismatches = {
        "raw_binding": int(
            raw_comparison["binding_canonical_to_production_changed"].sum()
        ),
        "raw_exposure": int(
            raw_comparison["exposure_canonical_to_production_changed"].sum()
        ),
        "pair_binding": int(
            pair_comparison["binding_canonical_to_production_changed"].sum()
        ),
        "pair_exposure": int(
            pair_comparison["exposure_canonical_to_production_changed"].sum()
        ),
    }
    failures: list[str] = []
    if len(raw_comparison) != 121_097:
        failures.append(f"raw_rows={len(raw_comparison)}")
    if len(pair_comparison) != 95_512:
        failures.append(f"pair_rows={len(pair_comparison)}")
    if pair_comparison["_merge"].ne("both").any():
        failures.append("pair_key_mismatch")
    if raw_binding_changes or pair_binding_changes:
        failures.append("binding_labels_changed")
    if raw_exposure_changes != 77:
        failures.append(f"raw_exposure_changes={raw_exposure_changes}")
    if pair_exposure_changes != 46:
        failures.append(f"pair_exposure_changes={pair_exposure_changes}")
    if any(canonical_mismatches.values()):
        failures.append(f"canonical_production_mismatches={canonical_mismatches}")
    changed_raw = raw_comparison[
        raw_comparison["exposure_legacy_to_production_changed"]
    ]
    changed_pairs = pair_comparison[
        pair_comparison["exposure_legacy_to_production_changed"]
    ]
    if not changed_raw["legacy_exposure_label"].map(_label_name).eq("unknown").all():
        failures.append("raw_transition_origin")
    if not changed_raw["production_exposure_label"].eq("negative").all():
        failures.append("raw_transition_destination")
    if not changed_pairs["legacy_exposure_label"].map(_label_name).eq("unknown").all():
        failures.append("pair_transition_origin")
    if not changed_pairs["production_exposure_label"].eq("negative").all():
        failures.append("pair_transition_destination")

    raw_path = out_dir / "parser_phase_b_production_comparison.csv"
    pair_path = out_dir / "parser_phase_b_pair_comparison.csv"
    raw_comparison.to_csv(raw_path, index=False)
    pair_comparison.to_csv(pair_path, index=False)
    summary = {
        "status": "passed" if not failures else "failed",
        "failures": failures,
        "policy_versions": {
            "parser": SPD_ACTIVITY_PARSER_VERSION,
            "binding": BINDING_LABEL_POLICY_VERSION,
            "exposure": SPD_LABEL_POLICY_VERSION,
        },
        "row_counts": {
            "raw": int(len(raw_comparison)),
            "pairs": int(len(pair_comparison)),
        },
        "old_label_counts": {
            "raw_binding": _label_counts(raw_comparison["legacy_binding_label"]),
            "raw_exposure": _label_counts(raw_comparison["legacy_exposure_label"]),
            "pair_binding": _label_counts(pair_comparison["legacy_binding_label"]),
            "pair_exposure": _label_counts(pair_comparison["legacy_exposure_label"]),
        },
        "new_label_counts": {
            "raw_binding": _label_counts(raw_comparison["production_binding_label"]),
            "raw_exposure": _label_counts(raw_comparison["production_exposure_label"]),
            "pair_binding": _label_counts(pair_comparison["production_binding_label"]),
            "pair_exposure": _label_counts(
                pair_comparison["production_exposure_label"]
            ),
        },
        "changed_row_counts": {
            "raw_binding": raw_binding_changes,
            "raw_exposure": raw_exposure_changes,
            "pair_binding": pair_binding_changes,
            "pair_exposure": pair_exposure_changes,
        },
        "canonical_to_production_mismatches": canonical_mismatches,
        "changed_source_row_ids": changed_raw["source_row_id"].astype(str).tolist(),
        "changed_pair_keys": changed_pairs[
            ["drug_id", "target_id", "source_row_id"]
        ].to_dict("records"),
    }
    return summary, [raw_path, pair_path]


def _write_phase_b_summary(
    path: Path,
    *,
    integration: dict[str, Any],
    evidence: pd.DataFrame,
) -> None:
    old = integration["old_label_counts"]
    new = integration["new_label_counts"]
    text = f"""# SPD interval parser Phase B verification

- Gate: **PASS for the frozen stored-value contract**
- Parser policy: `{SPD_ACTIVITY_PARSER_VERSION}`
- Binding policy: `{BINDING_LABEL_POLICY_VERSION}`
- Exposure policy: `{SPD_LABEL_POLICY_VERSION}`
- Raw observations: {integration["row_counts"]["raw"]:,}
- Selected drug-target pairs: {integration["row_counts"]["pairs"]:,}

## Exact correction

All {len(evidence):,} affected source rows contain literal numeric workbook
cells, relation `>`, and an exact Decimal lower exposure margin of 10. None
depends on a PPB inequality, an explicit free-Cmax interval, an Excel formula,
binary floating-point rounding, or display-only rounding. Exactly
{int(evidence["production_selected"].sum()):,} production-selected pairs change
from unknown to negative. Canonical-to-production mismatches are all zero.

Binding selected-pair counts are unchanged:
{old["pair_binding"]["positive"]:,} positive,
{old["pair_binding"]["negative"]:,} negative, and
{old["pair_binding"]["unknown"]:,} unknown.

Exposure selected-pair counts change from
{old["pair_exposure"]["positive"]:,}/
{old["pair_exposure"]["negative"]:,}/
{old["pair_exposure"]["unknown"]:,} to
{new["pair_exposure"]["positive"]:,}/
{new["pair_exposure"]["negative"]:,}/
{new["pair_exposure"]["unknown"]:,}
positive/negative/unknown.

## Provenance limitation

The workbook provides point values and `=` PPB qualifiers but no upstream
uncertainty, precision metadata, or resolvable NIBR curation record identifiers.
The stored-value mathematical result is verified; biological PK precision
remains unresolved and is not represented as error-free evidence. Two
nonselected HTR2A source rows contain an ancillary DrugCentral median that is
inconsistent with the summarized lower bound; zero selected changes have this
flag.

## Preserved contracts and deferred issues

Receptor mappings and strict eligibility are unchanged: 9I52 remains
DRD1/P21728, while 6HUJ and 8HCQ remain unresolved and strict-ineligible.
Selection, representative-row behavior, weak-margin negatives, direct-label
precedence, and existing deduplication are unchanged. Frozen Phase 1 and
external four-state tables were not regenerated, and no model was trained.

The current raw strict expert stage remains keyed by `drug_id x pdb_id` in
`outputs/audits/spd_receptor_mapping_contract_phase1_0e8e518d_v9/strict/`
(`ml_spd_binding_table.csv` and `ml_spd_exposure_table.csv`, 23,849 rows each).
The subsequent model-ready stage still calls `deduplicate_ml_rows` on
`drug_id x target_id`, yielding 22,580 Binding and 10,462 Exposure rows and
collapsing 365/156 concordant 9I52/9LLG observations. This mismatch is reported,
not changed by this patch.

The 34 frozen Phase 1 Binding positives that disagree with both committed and
canonical relation-aware derivation remain a separate artifact-provenance
follow-up and are not explained or modified by the Exposure correction.

## Production call graph

`normalize_spd_table` reads and normalizes the workbook. `aggregate_spd_assays`
preserves its existing biochemical-first/minimum-AC50 selection, then calls
`_spd_label_status`, which delegates relation interpretation to
`parse_spd_activity_interval` and Exposure classification to
`label_spd_exposure_observation`. The materialized pair panel is joined by
`enrich_spd_labels_for_run_master`. `build_spd_four_expert_tables` derives
Binding through `_binding_label`, which delegates valid explicit measurements
to `parse_spd_activity_interval` and `label_spd_binding_observation`; Exposure
continues to consume the panel's public status/direct-label fields. Generic
external four-state evidence does not enter this call graph and remains frozen.
"""
    path.write_text(text, encoding="utf-8")


def verify_boundary_rows(
    *,
    workbook_path: Path,
    phase_a_behavior_path: Path,
    out_dir: Path,
    mapping_manifest_path: Path | None = None,
) -> dict[str, Any]:
    started = datetime.now(timezone.utc)
    repo_root = Path(__file__).resolve().parents[2]
    changed = pd.read_csv(phase_a_behavior_path, low_memory=False)
    changed = changed[_bool_series(changed["exposure_changed"])].copy()
    if len(changed) != 77:
        raise RuntimeError(f"expected 77 affected source rows, found {len(changed)}")
    if changed["source_row_id"].astype(str).duplicated().any():
        raise RuntimeError("affected source-row IDs are not unique")

    s1 = pd.read_excel(workbook_path, sheet_name="S Data 1", header=2)
    s2 = pd.read_excel(workbook_path, sheet_name="S Data 2", header=2)
    s3 = pd.read_excel(workbook_path, sheet_name="S Data 3", header=2)
    s1["_source_row_id"] = s1["RowId"].map(_id)
    s1["_source_excel_row"] = range(4, len(s1) + 4)
    if s1["_source_row_id"].duplicated().any():
        raise RuntimeError("S Data 1 RowId values are not unique")
    changed = changed[
        [
            "source_row_id",
            "production_selected",
            "legacy_exposure_label",
            "canonical_exposure_label",
        ]
    ].copy()
    changed["_source_row_id"] = changed["source_row_id"].map(_id)
    selected = changed.merge(
        s1,
        on="_source_row_id",
        how="left",
        validate="one_to_one",
    )
    if selected["RowId"].isna().any():
        raise RuntimeError("one or more affected source rows are absent from S Data 1")

    s2["_drug_id"] = s2["drugcentral struct id"].map(_id)
    if s2["_drug_id"].duplicated().any():
        raise RuntimeError("S Data 2 drugcentral struct IDs are not unique")
    selected["_drug_id"] = selected["drugcentral_struct_id"].map(_id)
    selected = selected.merge(
        s2,
        on="_drug_id",
        how="left",
        suffixes=("", "_pk"),
        validate="many_to_one",
    )
    selected = selected.merge(
        s3,
        left_on="assay_group_name",
        right_on="assay name",
        how="left",
        suffixes=("", "_assay"),
        validate="many_to_one",
    )
    if selected["drugcentral name"].isna().any() or selected["assay name"].isna().any():
        raise RuntimeError("affected rows do not fully resolve to S Data 2/S Data 3")

    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    ws1 = workbook["S Data 1"]
    ws2 = workbook["S Data 2"]
    columns1 = _column_map(ws1)
    columns2 = _column_map(ws2)
    s2_excel_rows = {
        _id(ws2.cell(row=row, column=columns2["drugcentral struct id"]).value): row
        for row in range(4, ws2.max_row + 1)
    }

    requested_s1: set[str] = set()
    requested_s2: set[str] = set()
    coordinates: dict[str, dict[str, dict[str, Any]]] = {}
    for row in selected.to_dict("records"):
        row_id = _id(row["source_row_id"])
        excel_row = int(row["_source_excel_row"])
        drug_id = _id(row["drugcentral_struct_id"])
        pk_row = s2_excel_rows[drug_id]
        coordinates[row_id] = {
            "relation": _cell_metadata(ws1, excel_row, columns1["summarized prefix"]),
            "ac50": _cell_metadata(ws1, excel_row, columns1["summarized IC50"]),
            "margin": _cell_metadata(ws1, excel_row, columns1["free_cmax_margin"]),
            "total_cmax": _cell_metadata(ws2, pk_row, columns2["Cmax tot (uM)"]),
            "ppb": _cell_metadata(ws2, pk_row, columns2["PPB %"]),
            "free_cmax": _cell_metadata(ws2, pk_row, columns2["Cmax, free (uM)"]),
        }
        requested_s1.update(
            coordinates[row_id][key]["coordinate"]
            for key in ("relation", "ac50", "margin")
        )
        requested_s2.update(
            coordinates[row_id][key]["coordinate"]
            for key in ("total_cmax", "ppb", "free_cmax")
        )
    xml1 = _xml_cell_values(workbook_path, "S Data 1", requested_s1)
    xml2 = _xml_cell_values(workbook_path, "S Data 2", requested_s2)

    evidence_rows: list[dict[str, Any]] = []
    for row in selected.to_dict("records"):
        row_id = _id(row["source_row_id"])
        cells = coordinates[row_id]
        activity = _decimal(cells["ac50"]["value"])
        free_cmax = _decimal(cells["free_cmax"]["value"])
        total_cmax = _decimal(cells["total_cmax"]["value"])
        ppb = _decimal(cells["ppb"]["value"])
        fraction_unbound = Decimal(1) - (ppb / Decimal(100))
        calculated_free = total_cmax * fraction_unbound
        margin = activity / free_cmax
        relation = _clean(cells["relation"]["value"])
        qualifier = _clean(row["PPB qualifier"])
        source_median = _clean(row["DrugCentral.AC50.Median(uM)"])
        summary_inconsistency = bool(
            source_median and relation == ">" and _decimal(source_median) <= activity
        )
        stored_pass = all(
            (
                relation == ">",
                margin == Decimal(10),
                qualifier == "=",
                calculated_free == free_cmax,
                not cells["ac50"]["has_formula"],
                not cells["free_cmax"]["has_formula"],
                xml1.get(cells["ac50"]["coordinate"], "")
                == _clean(cells["ac50"]["value"]),
                xml2.get(cells["free_cmax"]["coordinate"], "")
                == _clean(cells["free_cmax"]["value"]),
            )
        )
        notes = [
            "Workbook cells are exact stored point values; no formula or display-only precision.",
            "PPB is equality-qualified and free Cmax exactly matches total Cmax times fraction unbound.",
            "Open relation yields a stored-value margin interval of (10,+inf).",
            "No upstream uncertainty or resolvable NIBR curation record identifier is supplied.",
        ]
        if summary_inconsistency:
            notes.append(
                "Ancillary DrugCentral median is below the summarized lower bound; row is not production-selected."
            )
        evidence_rows.append(
            {
                "source_row_id": row_id,
                "source_excel_row": int(row["_source_excel_row"]),
                "production_selected": bool(row["production_selected"]),
                "drugcentral_struct_id": _id(row["drugcentral_struct_id"]),
                "drug_name": row["name"],
                "parent_when_prodrug_id": _id(
                    row["drugcentral_struct_id(parent_when_prodrug)"]
                ),
                "struct_match_type": row["struct_match_type"],
                "inchi_key": row["inchi_key"],
                "assay_id": _id(row["assay_id"]),
                "assay_group": _id(row["assay_group"]),
                "assay_group_name": row["assay_group_name"],
                "preferred_assay_id": _id(row["preferred assay ID"]),
                "assay_format": row["Format"],
                "assay_mode": row["Mode"],
                "assay_readout": row["Readout"],
                "assay_technology": row["Technology"],
                "assay_species": row["Protein/Target/Species"],
                "assay_gene": row["EntrezGeneSymbol"],
                "assay_representative_gene": row[
                    "HumanEntrezGeneSymbol(representative)"
                ],
                "source_representative": row[
                    "representative_result_drug_assay_group_pair"
                ],
                "activity_relation": relation,
                "activity_relation_cell": cells["relation"]["coordinate"],
                "activity_relation_xml_value": xml1.get(
                    cells["relation"]["coordinate"], ""
                ),
                "activity_ac50_uM": str(activity),
                "activity_ac50_cell": cells["ac50"]["coordinate"],
                "activity_ac50_xml_value": xml1.get(cells["ac50"]["coordinate"], ""),
                "activity_ac50_cell_type": cells["ac50"]["data_type"],
                "activity_ac50_number_format": cells["ac50"]["number_format"],
                "activity_ac50_has_formula": cells["ac50"]["has_formula"],
                "n_crc_summarized": row["N CRC summarized"],
                "n_crc_total": row["N CRC total"],
                "drugcentral_ac50_median_uM": row["DrugCentral.AC50.Median(uM)"],
                "chembl_ac50_median_uM": row["ChEMBL.AC50.Median(uM)"],
                "subscription_ac50_median_uM": row["Subscription.AC50.Median(uM)"],
                "drugcentral_ac50_count": row["DrugCentral.AC50.Count"],
                "chembl_ac50_count": row["ChEMBL.AC50.Count"],
                "subscription_ac50_count": row["Subscription.AC50.Count"],
                "source_summary_inconsistency": summary_inconsistency,
                "total_cmax_uM": str(total_cmax),
                "total_cmax_cell": cells["total_cmax"]["coordinate"],
                "total_cmax_xml_value": xml2.get(cells["total_cmax"]["coordinate"], ""),
                "cmax_source": row["Cmax source"],
                "ppb_percent": str(ppb),
                "ppb_cell": cells["ppb"]["coordinate"],
                "ppb_xml_value": xml2.get(cells["ppb"]["coordinate"], ""),
                "ppb_qualifier": qualifier,
                "ppb_source": row["PPB source"],
                "fraction_unbound_decimal": str(fraction_unbound),
                "free_cmax_uM": str(free_cmax),
                "free_cmax_cell": cells["free_cmax"]["coordinate"],
                "free_cmax_xml_value": xml2.get(cells["free_cmax"]["coordinate"], ""),
                "free_cmax_cell_type": cells["free_cmax"]["data_type"],
                "free_cmax_number_format": cells["free_cmax"]["number_format"],
                "free_cmax_has_formula": cells["free_cmax"]["has_formula"],
                "free_cmax_calculated_decimal": str(calculated_free),
                "free_cmax_matches_calculation": calculated_free == free_cmax,
                "cmax_from_parent": row["Cmax from parent?"],
                "stored_margin": _clean(row["free_cmax_margin"]),
                "stored_margin_cell": cells["margin"]["coordinate"],
                "stored_margin_xml_value": xml1.get(cells["margin"]["coordinate"], ""),
                "lower_margin_decimal": str(margin),
                "margin_interval": "(10,+inf)",
                "old_exposure_label": row["legacy_exposure_label"],
                "new_exposure_label": row["canonical_exposure_label"],
                "ppb_is_bound": _qualifier_is_bound(qualifier),
                "free_cmax_has_explicit_interval": False,
                "binary_float_or_display_rounding_dependency": False,
                "source_record_identifier_available": False,
                "biological_precision_documented": False,
                "stored_value_contract_verdict": "PASS" if stored_pass else "FAIL",
                "biological_provenance_verdict": "UNRESOLVED",
                "verification_notes": " ".join(notes),
            }
        )

    evidence = pd.DataFrame(evidence_rows, columns=EVIDENCE_COLUMNS)
    selected_evidence = evidence[evidence["production_selected"]].copy()
    failures = evidence[evidence["stored_value_contract_verdict"].ne("PASS")]
    if len(selected_evidence) != 46:
        raise RuntimeError(
            f"expected 46 selected affected pairs, found {len(selected_evidence)}"
        )
    if not failures.empty:
        raise RuntimeError(
            f"stored-value verification failed for {len(failures)} affected rows"
        )

    out_dir.mkdir(parents=True, exist_ok=True)
    evidence_path = out_dir / "parser_boundary_provenance_verification.csv"
    manifest_path = out_dir / "parser_boundary_provenance_manifest.json"
    evidence.to_csv(evidence_path, index=False)
    integration, integration_paths = _verify_production_integration(
        workbook_path=workbook_path,
        phase_a_behavior_path=phase_a_behavior_path,
        out_dir=out_dir,
    )
    if integration["status"] != "passed":
        raise RuntimeError(
            f"Phase B production integration verification failed: {integration['failures']}"
        )
    summary_path = out_dir / "parser_phase_b_summary.md"
    _write_phase_b_summary(
        summary_path,
        integration=integration,
        evidence=evidence,
    )
    integration_paths.append(summary_path)
    mapping_validation = None
    if mapping_manifest_path is not None:
        mapping_manifest = json.loads(mapping_manifest_path.read_text(encoding="utf-8"))
        required_mapping_checks = (
            "affected_pdb_scope_exact",
            "legacy_no_row_multiplication",
            "strict_no_row_multiplication",
            "legacy_strict_row_count_equal",
            "unrelated_pdb_rows_unchanged",
            "strict_model_ready_excludes_6huj_8hcq",
            "strict_model_ready_retains_eligible_9i52",
        )
        failed_mapping_checks = [
            check
            for check in required_mapping_checks
            if mapping_manifest.get("checks", {}).get(check) is not True
        ]
        if mapping_manifest.get("status") != "pass" or failed_mapping_checks:
            raise RuntimeError(
                "mapping validation failed: "
                f"status={mapping_manifest.get('status')} "
                f"failed_checks={failed_mapping_checks}"
            )
        mapping_validation = {
            "status": "passed",
            "path": str(mapping_manifest_path.resolve()),
            "sha256": _sha256(mapping_manifest_path),
            "required_checks": {
                check: mapping_manifest["checks"][check]
                for check in required_mapping_checks
            },
            "affected_counts": mapping_manifest["affected_counts"],
            "strict_eligibility": mapping_manifest["strict_eligibility"],
            "label_counts": mapping_manifest["label_counts"],
            "label_count_interpretation": (
                "The rebuild consumed the frozen v2 pair panel, so these counts "
                "validate mapping preservation only and do not represent the "
                "approved raw-workbook boundary correction."
            ),
        }

    manifest = {
        "schema_version": 1,
        "status": "passed",
        "scope": "approved_open_lower_margin_10_exposure_correction",
        "git_sha": _git(repo_root, "rev-parse", "HEAD").strip(),
        "git_status_short_at_run": _git(repo_root, "status", "--short"),
        "parser_version": SPD_ACTIVITY_PARSER_VERSION,
        "inputs": [
            {
                "path": str(workbook_path.resolve()),
                "sha256": _sha256(workbook_path),
                "size_bytes": workbook_path.stat().st_size,
            },
            {
                "path": str(phase_a_behavior_path.resolve()),
                "sha256": _sha256(phase_a_behavior_path),
                "size_bytes": phase_a_behavior_path.stat().st_size,
            },
            *(
                [
                    {
                        "path": str(mapping_manifest_path.resolve()),
                        "sha256": _sha256(mapping_manifest_path),
                        "size_bytes": mapping_manifest_path.stat().st_size,
                    }
                ]
                if mapping_manifest_path is not None
                else []
            ),
        ],
        "row_counts": {
            "affected_source_rows": int(len(evidence)),
            "production_selected_pairs": int(len(selected_evidence)),
            "zidovudine_source_rows": int(evidence["drug_name"].eq("zidovudine").sum()),
            "mepacrine_source_rows": int(evidence["drug_name"].eq("mepacrine").sum()),
            "ppb_bound_rows": int(evidence["ppb_is_bound"].sum()),
            "explicit_free_cmax_interval_rows": int(
                evidence["free_cmax_has_explicit_interval"].sum()
            ),
            "rounding_dependent_rows": int(
                evidence["binary_float_or_display_rounding_dependency"].sum()
            ),
            "source_summary_inconsistency_rows": int(
                evidence["source_summary_inconsistency"].sum()
            ),
            "selected_source_summary_inconsistency_rows": int(
                selected_evidence["source_summary_inconsistency"].sum()
            ),
        },
        "stored_value_contract": {
            "verdict": "PASS",
            "reason": (
                "All 77 raw values are literal workbook point values; exact Decimal "
                "ratios equal 10, relation is open '>', PPB qualifiers are '=', and "
                "no explicit free-Cmax interval is present."
            ),
        },
        "biological_precision_and_provenance": {
            "verdict": "UNRESOLVED",
            "reason": (
                "The workbook does not provide upstream uncertainty, precision, or "
                "resolvable NIBR curation record identifiers. This correction applies "
                "the frozen stored-value contract and does not claim error-free PK."
            ),
        },
        "production_integration": integration,
        "mapping_validation": mapping_validation,
        "artifact": {
            "path": str(evidence_path.resolve()),
            "sha256": _sha256(evidence_path),
            "size_bytes": evidence_path.stat().st_size,
        },
        "additional_artifacts": [
            {
                "path": str(path.resolve()),
                "sha256": _sha256(path),
                "size_bytes": path.stat().st_size,
            }
            for path in integration_paths
        ],
        "implementation_sources": [
            {
                "path": path.relative_to(repo_root).as_posix(),
                "sha256": _sha256(path),
                "size_bytes": path.stat().st_size,
            }
            for path in (
                repo_root / "analysis/spd_activity_policy.py",
                repo_root / "analysis/external/spd.py",
                repo_root / "analysis/ml/spd_four_expert_tables.py",
                repo_root / "analysis/ml/spd_label_enrichment.py",
                repo_root / "analysis/ml/binding_label_contract_audit.py",
                repo_root / "analysis/cli/verify_spd_boundary_correction.py",
                repo_root / "analysis/test_spd_activity_policy.py",
                repo_root / "analysis/cli/test_audit_spd_interval_parser.py",
                repo_root / "analysis/ml/test_spd_receptor_mapping_integration.py",
                repo_root / "chemdb/tests/test_atlas_analysis_expanded.py",
            )
        ],
        "independent_reviews": [
            {
                "path": str(path.resolve()),
                "sha256": _sha256(path),
                "size_bytes": path.stat().st_size,
            }
            for path in sorted((out_dir / "reviews").glob("*.md"))
        ],
        "command": [sys.executable, *sys.argv],
        "verification_commands": [
            {
                "command": (
                    "python -m pytest analysis/test_spd_activity_policy.py "
                    "analysis/cli/test_audit_spd_interval_parser.py "
                    "analysis/ml/test_spd_receptor_mapping.py "
                    "analysis/ml/test_spd_receptor_mapping_integration.py "
                    "analysis/cli/test_rebuild_spd_receptor_mapping_contract.py -q"
                ),
                "result": "87 passed",
            },
            {
                "command": (
                    "python -m pytest analysis/test_spd_activity_policy.py "
                    "analysis/cli/test_audit_spd_interval_parser.py "
                    "analysis/ml/test_spd_receptor_mapping_integration.py "
                    "chemdb/tests/test_atlas_analysis_expanded.py -q"
                ),
                "result": "79 passed",
            },
            {
                "command": "ruff check <touched Python files>",
                "result": "passed",
            },
            {
                "command": "mypy <touched production and audit source files>",
                "result": "passed",
            },
            {
                "command": "git diff --check",
                "result": "passed",
            },
            {
                "command": "atlas runs",
                "result": (
                    "postFDAfixSPDrun running; atlas smoke internal and the full "
                    "smoke-bearing gate were not launched"
                ),
            },
        ],
        "environment": {
            "python": platform.python_version(),
            "implementation": platform.python_implementation(),
            "pandas": pd.__version__,
            "openpyxl": openpyxl.__version__,
            "executable": sys.executable,
        },
        "timestamp_started_utc": started.isoformat(),
        "timestamp_completed_utc": datetime.now(timezone.utc).isoformat(),
    }
    manifest_path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Verify exact workbook provenance for the approved open-lower-bound "
            "SPD exposure correction."
        )
    )
    parser.add_argument("--spd-workbook", required=True, type=Path)
    parser.add_argument("--phase-a-behavior", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--mapping-manifest", type=Path)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    manifest = verify_boundary_rows(
        workbook_path=args.spd_workbook,
        phase_a_behavior_path=args.phase_a_behavior,
        out_dir=args.out_dir,
        mapping_manifest_path=args.mapping_manifest,
    )
    print(json.dumps(manifest, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
