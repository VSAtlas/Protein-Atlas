from __future__ import annotations

import csv
import json
import logging
import re
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Iterable, Mapping, Optional, Sequence, cast

import requests  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]

from analysis.reporting.fda_name_map import (
    resolve_ligand_display_name,
    resolve_mapping_csv_path,
    try_load_fda_index,
)
from analysis.reporting.value_utils import normalize_side_effect_label
from analysis.reporting.target_safety_evidence import (
    SAFETY_DEFAULT,
    classify_safety_bucket_scores,
    safety_bucket_order,
)

OPENFDA_EVENT_URL = "https://api.fda.gov/drug/event.json"
OPENFDA_LABEL_URL = "https://api.fda.gov/drug/label.json"
OPENFDA_SOURCE = "openFDA FAERS"
OPENFDA_LABEL_SOURCE = "openFDA drug label"
SPD_EXPOSURE_SOURCE = "SPD Supplementary Data 2 (s41467-023-40064-9)"
LIGAND_SIDE_EFFECT_CACHE_VERSION = 1
DEFAULT_CACHE_REL = Path("pathways") / "cache" / "ligand_side_effects_openfda.json"

_LOG = logging.getLogger(__name__)
_RDK_RE = re.compile(r"rdk[_-]?(\d+)", re.IGNORECASE)
_SPLIT_RE = re.compile(r"[|;]")
_EXCLUDED_EVENT_FRAGMENTS = (
    "accidental exposure",
    "completed suicide",
    "death",
    "disease progression",
    "drug dependence",
    "drug ineffective",
    "drug interaction",
    "drug resistance",
    "intentional product misuse",
    "intentional product use issue",
    "lack of efficacy",
    "malignant neoplasm progression",
    "medication error",
    "no adverse event",
    "off label use",
    "product administration error",
    "product dose omission issue",
    "product quality issue",
    "suicide attempt",
    "therapeutic response decreased",
    "wrong technique in product usage process",
)
_OPENFDA_PHRASE_PUNCT_RE = re.compile(r'[+"?*~^{}\[\]\\():]')
_NAME_FIELDS = (
    "display_name",
    "rxnorm_generic_name",
    "drugcentral_generic_name",
    "generic_name",
    "pubchem_name",
    "pubchem_record_title",
    "sdf_title",
    "remark_name",
)
_QUERY_NAME_FIELDS = (
    "rxnorm_generic_name",
    "drugcentral_generic_name",
    "generic_name",
    "display_name",
    "pubchem_record_title",
    "remark_name",
)
_BRAND_FIELDS = ("brand_names", "rxnorm_brand_names", "drugcentral_brand_names")
_FORMULA_FIELDS = ("sdf_formula", "pdbqt_formula")
_ATOMIC_WEIGHTS = {
    "H": 1.00794,
    "C": 12.0107,
    "N": 14.0067,
    "O": 15.9994,
    "F": 18.9984032,
    "P": 30.973762,
    "S": 32.065,
    "Cl": 35.453,
    "Br": 79.904,
    "I": 126.90447,
    "Na": 22.98976928,
    "K": 39.0983,
    "Ca": 40.078,
    "Mg": 24.305,
    "Zn": 65.38,
}
_FORMULA_TOKEN_RE = re.compile(r"([A-Z][a-z]?)(\d*)")
_CMAX_VALUE_RE = re.compile(
    r"(?:(?:c\s*max|cmax|peak plasma concentration|maximum plasma concentration)"
    r"[^.;:\n]{0,180}?|[^.;:\n]{0,80}?"
    r"(?:c\s*max|cmax))"
    r"([0-9]+(?:\.[0-9]+)?)\s*"
    r"(ng/ml|ng/mL|mcg/ml|mcg/mL|ug/ml|ug/mL|µg/ml|µg/mL|mg/l|mg/L|nmol/l|nM|umol/l|µM|uM)",
    re.IGNORECASE,
)
_PROTEIN_BINDING_RE = re.compile(
    r"(?:protein binding|plasma protein binding|bound to (?:human )?plasma proteins)"
    r"[^.;:\n]{0,120}?"
    r"([<>]?\s*[0-9]+(?:\.[0-9]+)?)\s*%",
    re.IGNORECASE,
)
_LOW_INFORMATION_NAMES = {
    "empty",
    "n a",
    "na",
    "not applicable",
    "none",
    "null",
    "unknown",
}


@dataclass(frozen=True)
class LigandQuery:
    ligand_label: str
    ligand_base: str
    ligand_display: str
    query_names: tuple[str, ...]
    brand_names: tuple[str, ...]
    rxcuis: tuple[str, ...]
    uniis: tuple[str, ...]
    molecular_weight: float | None
    mapping_used: bool


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _normalize_key(value: Any) -> str:
    text = normalize_side_effect_label(value).casefold()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _is_placeholder_name(value: Any) -> bool:
    text = normalize_side_effect_label(value)
    if not text:
        return True
    lower = text.casefold()
    if lower in {"nan", "none", "null", "unknown", "unk"}:
        return True
    if re.fullmatch(r"fda[_ -]?\d+", lower):
        return True
    if re.fullmatch(r"rdk[_ -]?\d+", lower):
        return True
    return False


def _is_pdb_component_label(value: Any) -> bool:
    text = normalize_side_effect_label(value)
    return bool(re.fullmatch(r"[A-Za-z0-9]{2,4}\s+[A-Za-z]\d{2,5}", text))


def _looks_like_formula(value: Any) -> bool:
    text = str(value or "").strip()
    if not text:
        return False
    return bool(re.fullmatch(r"(?:[A-Z][a-z]?\d*){2,}", text))


def _is_query_name(value: Any) -> bool:
    text = normalize_side_effect_label(value)
    if _is_placeholder_name(text) or _is_pdb_component_label(text):
        return False
    normalized = _normalize_key(text)
    if normalized in _LOW_INFORMATION_NAMES:
        return False
    compact = re.sub(r"[^A-Za-z0-9]+", "", text)
    alpha_count = sum(1 for char in compact if char.isalpha())
    if alpha_count < 5:
        return False
    if re.fullmatch(r"[A-Za-z]{1,5}", compact):
        return False
    if re.fullmatch(r"(?:cid|unii|ncgc)[\s_-]*[A-Za-z0-9-]+", text, re.IGNORECASE):
        return False
    if re.fullmatch(r"\d+(?:[\s-]\d+){1,}", text):
        return False
    return not _looks_like_formula(text)


def _dedupe(values: Iterable[Any], *, limit: int = 0) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        clean = normalize_side_effect_label(value)
        key = clean.casefold()
        if not clean or key in seen:
            continue
        seen.add(key)
        out.append(clean)
        if limit and len(out) >= limit:
            break
    return out


def _dedupe_preserve(values: Iterable[str], *, limit: int = 0) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        clean = str(value or "").strip()
        key = clean.casefold()
        if not clean or key in seen:
            continue
        seen.add(key)
        out.append(clean)
        if limit and len(out) >= limit:
            break
    return out


def _molecular_weight_from_formula(value: Any) -> float | None:
    formula = normalize_side_effect_label(value)
    if not formula:
        return None
    weight = 0.0
    consumed = ""
    for match in _FORMULA_TOKEN_RE.finditer(formula):
        element = match.group(1)
        count = int(match.group(2) or "1")
        atomic_weight = _ATOMIC_WEIGHTS.get(element)
        if atomic_weight is None:
            return None
        weight += atomic_weight * count
        consumed += match.group(0)
    if not consumed or consumed != formula:
        return None
    return weight if weight > 0 else None


def _molecular_weight_from_mapping(mapping_row: Mapping[str, Any]) -> float | None:
    for field in _FORMULA_FIELDS:
        weight = _molecular_weight_from_formula(mapping_row.get(field))
        if weight is not None:
            return weight
    return None


def _split_values(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [part.strip() for part in _SPLIT_RE.split(value) if part.strip()]
    if isinstance(value, Sequence) and not isinstance(value, (bytes, bytearray)):
        out: list[str] = []
        for part in value:
            clean = normalize_side_effect_label(part)
            if clean:
                out.append(clean)
        return out
    text = normalize_side_effect_label(value)
    return [text] if text else []


def _extract_rdk_id(value: Any) -> str:
    match = _RDK_RE.search(str(value or ""))
    if not match:
        return ""
    return f"rdk_{match.group(1).zfill(7)}"


def _as_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _spd_fraction_unbound(ppb_percent: Any) -> float | None:
    bound_pct = _as_float(ppb_percent)
    if bound_pct is None:
        return None
    return max(0.0, min(1.0, 1.0 - (bound_pct / 100.0)))


def cache_path(repo_root: Path) -> Path:
    return repo_root / DEFAULT_CACHE_REL


@lru_cache(maxsize=4)
def _load_cache(path_value: str) -> Dict[str, Any]:
    path = Path(path_value)
    if not path.exists():
        return {
            "version": LIGAND_SIDE_EFFECT_CACHE_VERSION,
            "source": OPENFDA_SOURCE,
            "entries": {},
            "by_key": {},
        }
    payload = json.loads(path.read_text(encoding="utf-8"))
    return cast(Dict[str, Any], payload if isinstance(payload, dict) else {})


def clear_ligand_side_effect_cache() -> None:
    _load_cache.cache_clear()


def load_ligand_side_effect_cache(repo_root: Path) -> Dict[str, Any]:
    return _load_cache(str(cache_path(repo_root)))


def _cache_lookup_keys(
    ligand_label: str, ligand_base: str = "", ligand_display: str = ""
) -> list[str]:
    return _dedupe(
        [
            ligand_base,
            _extract_rdk_id(ligand_base),
            _extract_rdk_id(ligand_label),
            ligand_label,
            ligand_display,
        ]
    )


def lookup_ligand_side_effect_entry(
    repo_root: Path,
    *,
    ligand_label: str,
    ligand_base: str = "",
    ligand_display: str = "",
) -> Dict[str, Any]:
    payload = load_ligand_side_effect_cache(repo_root)
    entries = cast(Dict[str, Any], payload.get("entries") or {})
    by_key = cast(Dict[str, Any], payload.get("by_key") or {})
    for raw_key in _cache_lookup_keys(ligand_label, ligand_base, ligand_display):
        normalized = _normalize_key(raw_key)
        entry_key = str(by_key.get(normalized) or "")
        if entry_key and isinstance(entries.get(entry_key), dict):
            return cast(Dict[str, Any], entries[entry_key])
        if isinstance(entries.get(normalized), dict):
            return cast(Dict[str, Any], entries[normalized])
    return {}


def load_fda_mapping_rows(mapping_csv: Optional[Path]) -> Dict[str, Dict[str, Any]]:
    if mapping_csv is None or not mapping_csv.exists():
        return {}
    out: Dict[str, Dict[str, Any]] = {}
    with mapping_csv.open("r", encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            lowered = {str(k or "").strip().lower(): v for k, v in row.items()}
            keys: list[str] = []
            rdk_id = _extract_rdk_id(lowered.get("rdk_id")) or _extract_rdk_id(
                lowered.get("path")
            )
            if not rdk_id and str(lowered.get("scheme") or "").strip().lower() == "rdk":
                try:
                    rdk_id = f"rdk_{int(float(str(lowered.get('file_num') or ''))):07d}"
                except ValueError:
                    rdk_id = ""
            if rdk_id:
                keys.append(rdk_id)
            for field in (*_NAME_FIELDS, *_BRAND_FIELDS):
                keys.extend(_split_values(lowered.get(field)))
            for key in keys:
                normalized = _normalize_key(key)
                if normalized and normalized not in out:
                    out[normalized] = lowered
    return out


def mapping_row_for_ligand(
    mapping_rows: Mapping[str, Mapping[str, Any]],
    *,
    ligand_base: str,
    ligand_display: str,
) -> Mapping[str, Any]:
    for key in _cache_lookup_keys(ligand_display, ligand_base, ligand_display):
        row = mapping_rows.get(_normalize_key(key))
        if row:
            return row
    return {}


def build_ligand_query(
    repo_root: Path,
    row: Mapping[str, Any],
    *,
    mapping_csv: Optional[Path] = None,
    fda_index: Any = None,
    mapping_rows: Optional[Mapping[str, Mapping[str, Any]]] = None,
) -> LigandQuery:
    ligand_base = normalize_side_effect_label(row.get("ligand_base"))
    ligand_label = normalize_side_effect_label(row.get("ligand_name")) or normalize_side_effect_label(
        row.get("ligand_display")
    ) or ligand_base
    ligand_display = normalize_side_effect_label(row.get("ligand_display")) or ligand_label
    mapping_used = False
    if _is_placeholder_name(ligand_display) and fda_index is not None:
        resolved = normalize_side_effect_label(
            resolve_ligand_display_name(ligand_base, ligand_base, fda_index)
        )
        if resolved and not _is_placeholder_name(resolved):
            ligand_display = resolved
            mapping_used = True

    rows_by_key = mapping_rows or load_fda_mapping_rows(mapping_csv)
    mapping_row = mapping_row_for_ligand(
        rows_by_key, ligand_base=ligand_base, ligand_display=ligand_display
    )
    mapping_used = mapping_used or bool(mapping_row)
    names: list[str] = []
    if mapping_row:
        for field in _QUERY_NAME_FIELDS:
            names.extend(_split_values(mapping_row.get(field)))
    else:
        names.extend([ligand_display, ligand_label])
    names = [name for name in names if _is_query_name(name)]
    brands: list[str] = []
    for field in _BRAND_FIELDS:
        brands.extend(_split_values(mapping_row.get(field)))
    return LigandQuery(
        ligand_label=ligand_label,
        ligand_base=ligand_base,
        ligand_display=ligand_display,
        query_names=tuple(_dedupe(names, limit=8)),
        brand_names=tuple(
            _dedupe([brand for brand in brands if _is_query_name(brand)], limit=10)
        ),
        rxcuis=tuple(_dedupe(_split_values(mapping_row.get("rxnorm_rxcui")), limit=4)),
        uniis=tuple(_dedupe(_split_values(mapping_row.get("pubchem_unii_list")), limit=4)),
        molecular_weight=_molecular_weight_from_mapping(mapping_row),
        mapping_used=mapping_used,
    )


def _openfda_event_terms(
    session: requests.Session,
    search: str,
    *,
    limit: int,
    timeout: int,
) -> tuple[list[Dict[str, Any]], str]:
    response = _openfda_get(
        session,
        OPENFDA_EVENT_URL,
        params={
            "search": search,
            "count": "patient.reaction.reactionmeddrapt.exact",
            "limit": str(limit),
        },
        timeout=timeout,
    )
    if response.status_code == 404:
        return [], ""
    response.raise_for_status()
    payload = response.json()
    meta = cast(Dict[str, Any], payload.get("meta") or {})
    return cast(list[Dict[str, Any]], payload.get("results") or []), str(
        meta.get("last_updated") or ""
    )


def _openfda_label_records(
    session: requests.Session,
    search: str,
    *,
    limit: int,
    timeout: int,
) -> tuple[list[Dict[str, Any]], str]:
    response = _openfda_get(
        session,
        OPENFDA_LABEL_URL,
        params={
            "search": search,
            "limit": str(limit),
        },
        timeout=timeout,
    )
    if response.status_code == 404:
        return [], ""
    response.raise_for_status()
    payload = response.json()
    meta = cast(Dict[str, Any], payload.get("meta") or {})
    return cast(list[Dict[str, Any]], payload.get("results") or []), str(
        meta.get("last_updated") or ""
    )


def _openfda_get(
    session: requests.Session,
    url: str,
    *,
    params: Mapping[str, str],
    timeout: int,
    retries: int = 4,
) -> requests.Response:
    last_response: requests.Response | None = None
    for attempt in range(max(1, retries)):
        response = session.get(url, params=params, timeout=timeout)
        last_response = response
        if response.status_code not in {429, 500, 502, 503, 504}:
            return response
        retry_after = response.headers.get("Retry-After")
        try:
            delay = float(retry_after) if retry_after else 0.0
        except ValueError:
            delay = 0.0
        delay = max(delay, min(60.0, 2.0 * float(attempt + 1)))
        _LOG.warning(
            "openFDA request throttled status=%s attempt=%d/%d sleep=%.1fs",
            response.status_code,
            attempt + 1,
            retries,
            delay,
        )
        time.sleep(delay)
    if last_response is None:
        raise RuntimeError("openFDA request did not execute")
    return last_response


def _query_strings(query: LigandQuery) -> list[str]:
    searches: list[str] = []
    for rxcui in query.rxcuis:
        searches.append(_openfda_phrase_search("patient.drug.openfda.rxcui", rxcui))
    for unii in query.uniis:
        searches.append(_openfda_phrase_search("patient.drug.openfda.unii", unii))
    for brand in query.brand_names:
        searches.append(_openfda_phrase_search("patient.drug.openfda.brand_name", brand))
    for name in query.query_names:
        searches.append(_openfda_phrase_search("patient.drug.openfda.generic_name", name))
        searches.append(_openfda_phrase_search("patient.drug.openfda.brand_name", name))
    return _dedupe_preserve(searches, limit=40)


def _label_query_strings(query: LigandQuery) -> list[str]:
    searches: list[str] = []
    for rxcui in query.rxcuis:
        searches.append(_openfda_phrase_search("openfda.rxcui", rxcui))
    for unii in query.uniis:
        searches.append(_openfda_phrase_search("openfda.unii", unii))
    for brand in query.brand_names:
        searches.append(_openfda_phrase_search("openfda.brand_name", brand))
    for name in query.query_names:
        searches.append(_openfda_phrase_search("openfda.generic_name", name))
        searches.append(_openfda_phrase_search("openfda.brand_name", name))
    return _dedupe_preserve(searches, limit=40)


def _is_side_effect_term(term: str) -> bool:
    normalized = _normalize_key(term)
    if not normalized:
        return False
    return not any(fragment in normalized for fragment in _EXCLUDED_EVENT_FRAGMENTS)


def _format_term(value: Any) -> str:
    text = normalize_side_effect_label(value).lower()
    if not text:
        return ""
    return " ".join(part.capitalize() for part in text.split())


def _openfda_phrase_search(field: str, value: Any) -> str:
    clean = normalize_side_effect_label(value)
    clean = _OPENFDA_PHRASE_PUNCT_RE.sub(" ", clean)
    clean = " ".join(clean.split()).strip(" -")
    if not clean:
        return ""
    return f'{field}:"{clean}"'


def _flatten_label_text(record: Mapping[str, Any]) -> str:
    fields = (
        "pharmacokinetics",
        "clinical_pharmacology",
        "description",
        "pharmacodynamics",
    )
    parts: list[str] = []
    for field in fields:
        raw = record.get(field)
        if isinstance(raw, str):
            parts.append(raw)
        elif isinstance(raw, Sequence) and not isinstance(raw, (bytes, bytearray)):
            parts.extend(str(item) for item in raw if normalize_side_effect_label(item))
    return " ".join(" ".join(parts).split())


def _cmax_to_um(value: float, unit: str, molecular_weight: float | None) -> float | None:
    unit_norm = unit.replace("µ", "u").lower()
    if unit_norm in {"um", "umol/l"}:
        return value
    if unit_norm in {"nm", "nmol/l"}:
        return value / 1000.0
    if molecular_weight is None or molecular_weight <= 0:
        return None
    if unit_norm in {"ng/ml"}:
        mg_per_l = value * 0.001
    elif unit_norm in {"mcg/ml", "ug/ml", "mg/l"}:
        mg_per_l = value
    else:
        return None
    return (mg_per_l * 1000.0) / molecular_weight


def _extract_label_exposure(
    records: Sequence[Mapping[str, Any]],
    *,
    molecular_weight: float | None,
) -> Dict[str, Any]:
    cmax_um: float | None = None
    cmax_source_unit = ""
    fraction_unbound: float | None = None
    for record in records:
        text = _flatten_label_text(record)
        if not text:
            continue
        if cmax_um is None:
            for match in _CMAX_VALUE_RE.finditer(text):
                raw_value = float(match.group(1))
                raw_unit = match.group(2)
                converted = _cmax_to_um(raw_value, raw_unit, molecular_weight)
                if converted is not None and converted > 0:
                    cmax_um = converted
                    cmax_source_unit = f"{raw_value:g} {raw_unit}"
                    break
        if fraction_unbound is None:
            binding_match = _PROTEIN_BINDING_RE.search(text)
            if binding_match:
                raw_bound = binding_match.group(1).replace(" ", "")
                is_greater_than = raw_bound.startswith(">")
                bound_pct = float(raw_bound.lstrip("<>"))
                if is_greater_than and bound_pct >= 99.0:
                    fraction_unbound = 0.01
                else:
                    fraction_unbound = max(0.0, min(1.0, 1.0 - bound_pct / 100.0))
        if cmax_um is not None and fraction_unbound is not None:
            break

    out: Dict[str, Any] = {}
    if cmax_um is not None:
        out["cmax_um"] = round(cmax_um, 6)
    if fraction_unbound is not None:
        out["fraction_unbound_plasma"] = round(fraction_unbound, 6)
    if cmax_um is not None and fraction_unbound is not None:
        out["free_cmax_um"] = round(cmax_um * fraction_unbound, 6)
    if out:
        details = []
        if cmax_source_unit:
            details.append(f"Cmax {cmax_source_unit}")
        if molecular_weight is not None:
            details.append(f"MW {molecular_weight:.3f} g/mol")
        if fraction_unbound is not None:
            details.append(f"fu {fraction_unbound:.4g}")
        out["exposure_source"] = f"{OPENFDA_LABEL_SOURCE}: " + "; ".join(details)
    return out


def fetch_ligand_label_exposure_entry(
    query: LigandQuery,
    *,
    session: Optional[requests.Session] = None,
    label_limit: int = 5,
    timeout: int = 20,
    sleep_sec: float = 0.0,
) -> Dict[str, Any]:
    client = session or requests.Session()
    last_updated = ""
    selected_search = ""
    status = "no_query"
    records: list[Dict[str, Any]] = []
    for search in _label_query_strings(query):
        if sleep_sec > 0:
            time.sleep(sleep_sec)
        try:
            records, query_last_updated = _openfda_label_records(
                client, search, limit=label_limit, timeout=timeout
            )
        except Exception as exc:
            status = "query_failed"
            _LOG.warning("openFDA label query failed search=%s error=%s", search, exc)
            continue
        if query_last_updated:
            last_updated = query_last_updated
        if not records:
            status = "no_results"
            continue
        selected_search = search
        status = "ok"
        break

    out = _extract_label_exposure(records, molecular_weight=query.molecular_weight)
    out["openfda_label_search"] = selected_search
    out["openfda_label_last_updated"] = last_updated
    out["openfda_label_status"] = status
    out["openfda_label_url"] = OPENFDA_LABEL_URL
    return out


def _bucket_scores_for_terms(
    repo_root: Path, terms: Sequence[tuple[str, int]]
) -> Dict[str, float]:
    scores: Dict[str, float] = {bucket: 0.0 for bucket in safety_bucket_order(repo_root)}
    for term, count in terms:
        term_scores = classify_safety_bucket_scores(
            repo_root, event_text=term, evidence_source="drug_ae"
        )
        weight = min(3.0, 1.0 + (max(0, count).bit_length() / 12.0))
        for bucket, score in term_scores.items():
            scores[bucket] = scores.get(bucket, 0.0) + float(score) * weight
    return {bucket: round(score, 6) for bucket, score in scores.items() if score > 0.0}


def _bucket_fields(repo_root: Path, scores: Mapping[str, float]) -> Dict[str, Any]:
    order = safety_bucket_order(repo_root)
    buckets = [bucket for bucket in order if float(scores.get(bucket, 0.0)) > 0.0]
    if not buckets:
        return {
            "safety_buckets": [],
            "primary_display_safety": "",
            "secondary_safety_buckets": [],
            "drug_ae_buckets": [],
            "safety_bucket_scores": {},
        }
    primary = max(
        buckets,
        key=lambda bucket: (float(scores.get(bucket, 0.0)), -order.index(bucket)),
    )
    secondary = [bucket for bucket in buckets if bucket != primary]
    return {
        "safety_buckets": [primary, *secondary],
        "primary_display_safety": primary or SAFETY_DEFAULT,
        "secondary_safety_buckets": secondary,
        "drug_ae_buckets": buckets,
        "safety_bucket_scores": {
            bucket: f"{float(scores.get(bucket, 0.0)):.3f}" for bucket in buckets
        },
    }


def fetch_ligand_side_effect_entry(
    repo_root: Path,
    query: LigandQuery,
    *,
    session: Optional[requests.Session] = None,
    event_limit: int = 50,
    label_limit: int = 5,
    max_side_effects: int = 20,
    timeout: int = 20,
    sleep_sec: float = 0.0,
    fetch_label_exposure: bool = True,
) -> Dict[str, Any]:
    client = session or requests.Session()
    query_results: list[Dict[str, Any]] = []
    last_updated = ""
    selected_search = ""
    status = "no_query"
    for search in _query_strings(query):
        if sleep_sec > 0:
            time.sleep(sleep_sec)
        try:
            results, query_last_updated = _openfda_event_terms(
                client, search, limit=event_limit, timeout=timeout
            )
        except Exception as exc:
            status = "query_failed"
            _LOG.warning("openFDA ligand side-effect query failed search=%s error=%s", search, exc)
            continue
        if query_last_updated:
            last_updated = query_last_updated
        if not results:
            status = "no_results"
            continue
        query_results = results
        selected_search = search
        status = "ok"
        break

    accepted: list[tuple[str, int]] = []
    counts: Dict[str, int] = {}
    for item in query_results:
        term = _format_term(item.get("term"))
        if not _is_side_effect_term(term):
            continue
        count = int(float(str(item.get("count") or "0")))
        key = term.casefold()
        if key in counts:
            continue
        counts[key] = count
        accepted.append((term, count))
        if len(accepted) >= max_side_effects:
            break

    side_effects = [term for term, _count in accepted]
    bucket_scores = _bucket_scores_for_terms(repo_root, accepted)
    entry: Dict[str, Any] = {
        "ligand_label": query.ligand_label,
        "ligand_base": query.ligand_base,
        "ligand_display": query.ligand_display,
        "query_names": list(query.query_names),
        "brand_names": list(query.brand_names),
        "side_effects": side_effects,
        "side_effect_counts": {term: count for term, count in accepted},
        "side_effect_sources": [OPENFDA_SOURCE] if side_effects else [],
        "side_effect_evidence_summary": (
            f"{OPENFDA_SOURCE}: {len(side_effects)} reported reaction terms"
            if side_effects
            else ""
        ),
        "safety_confidence": "low" if side_effects else "unassigned",
        "openfda_event_search": selected_search,
        "openfda_last_updated": last_updated,
        "openfda_status": status,
        "openfda_url": OPENFDA_EVENT_URL,
        "mapping_used": query.mapping_used,
    }
    entry.update(_bucket_fields(repo_root, bucket_scores))
    if fetch_label_exposure:
        entry.update(
            fetch_ligand_label_exposure_entry(
                query,
                session=client,
                label_limit=label_limit,
                timeout=timeout,
                sleep_sec=sleep_sec,
            )
        )
    return entry


def build_cache_payload(entries: Sequence[Mapping[str, Any]]) -> Dict[str, Any]:
    payload_entries: Dict[str, Any] = {}
    by_key: Dict[str, str] = {}
    for entry in entries:
        primary = _normalize_key(
            entry.get("ligand_base")
            or entry.get("ligand_label")
            or entry.get("ligand_display")
        )
        if not primary:
            continue
        payload_entries[primary] = dict(entry)
        for raw_key in (
            entry.get("ligand_display"),
            entry.get("ligand_label"),
            entry.get("ligand_base"),
            *_split_values(entry.get("query_names")),
            *_split_values(entry.get("brand_names")),
        ):
            key = _normalize_key(raw_key)
            if key and key not in by_key:
                by_key[key] = primary
    return {
        "version": LIGAND_SIDE_EFFECT_CACHE_VERSION,
        "generated_at": _now_iso(),
        "source": OPENFDA_SOURCE,
        "source_url": OPENFDA_EVENT_URL,
        "entries": payload_entries,
        "by_key": by_key,
    }


def write_cache(repo_root: Path, payload: Mapping[str, Any]) -> Path:
    path = cache_path(repo_root)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    clear_ligand_side_effect_cache()
    return path


def resolve_mapping_context(repo_root: Path, run_id: str, fda_mapping_csv: str = "") -> tuple[Optional[Path], Any, Dict[str, Dict[str, Any]]]:
    mapping_csv = resolve_mapping_csv_path(
        repo_root, run_id, cli_value=fda_mapping_csv or None
    )
    fda_index = try_load_fda_index(mapping_csv)
    mapping_rows = load_fda_mapping_rows(mapping_csv)
    return mapping_csv, fda_index, mapping_rows


def load_spd_exposure_rows(spd_xlsx: Optional[Path]) -> Dict[str, Dict[str, Any]]:
    if spd_xlsx is None or not spd_xlsx.exists():
        return {}
    wb = load_workbook(spd_xlsx, read_only=True, data_only=True)
    try:
        if "S Data 2" not in wb.sheetnames:
            return {}
        ws = wb["S Data 2"]
        header_map: Dict[str, int] = {}
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if "drugcentral name" in [v.casefold() for v in values]:
                header_map = {str(v).strip().casefold(): idx for idx, v in enumerate(values) if v is not None}
                break
        if not header_map:
            return {}
        name_idx = header_map.get("drugcentral name")
        cmax_idx = header_map.get("cmax tot (um)") or header_map.get("cmax tot (uM)")
        ppb_idx = header_map.get("ppb %")
        free_idx = header_map.get("cmax, free (um)") or header_map.get("cmax, free (uM)")
        cmax_source_idx = header_map.get("cmax source")
        ppb_source_idx = header_map.get("ppb source")
        if name_idx is None:
            return {}
        out: Dict[str, Dict[str, Any]] = {}
        for row in ws.iter_rows(min_row=13, values_only=True):
            if name_idx >= len(row):
                continue
            raw_name = row[name_idx]
            if raw_name is None:
                continue
            name = normalize_side_effect_label(raw_name)
            key = _normalize_key(name)
            if not key:
                continue
            cmax_um = _as_float(row[cmax_idx]) if cmax_idx is not None and cmax_idx < len(row) else None
            free_cmax_um = _as_float(row[free_idx]) if free_idx is not None and free_idx < len(row) else None
            fraction_unbound = _spd_fraction_unbound(row[ppb_idx]) if ppb_idx is not None and ppb_idx < len(row) else None
            if free_cmax_um is None and cmax_um is not None and fraction_unbound is not None:
                free_cmax_um = cmax_um * fraction_unbound
            if cmax_um is None and free_cmax_um is None and fraction_unbound is None:
                continue
            source_parts = [SPD_EXPOSURE_SOURCE]
            if cmax_source_idx is not None and cmax_source_idx < len(row):
                cmax_src = normalize_side_effect_label(row[cmax_source_idx])
                if cmax_src:
                    source_parts.append(f"Cmax source: {cmax_src}")
            if ppb_source_idx is not None and ppb_source_idx < len(row):
                ppb_src = normalize_side_effect_label(row[ppb_source_idx])
                if ppb_src:
                    source_parts.append(f"PPB source: {ppb_src}")
            payload: Dict[str, Any] = {"exposure_source": "; ".join(source_parts)}
            if cmax_um is not None and cmax_um > 0:
                payload["cmax_um"] = round(cmax_um, 6)
            if fraction_unbound is not None:
                payload["fraction_unbound_plasma"] = round(fraction_unbound, 6)
            if free_cmax_um is not None and free_cmax_um > 0:
                payload["free_cmax_um"] = round(free_cmax_um, 6)
            out[key] = payload
        return out
    finally:
        wb.close()


def spd_exposure_for_query(
    spd_rows: Mapping[str, Mapping[str, Any]],
    *,
    query: LigandQuery,
) -> Dict[str, Any]:
    if not spd_rows:
        return {}
    candidates = _dedupe(
        [
            query.ligand_display,
            query.ligand_label,
            query.ligand_base,
            *query.query_names,
            *query.brand_names,
        ]
    )
    for candidate in candidates:
        match = spd_rows.get(_normalize_key(candidate))
        if match:
            return dict(match)
    return {}
